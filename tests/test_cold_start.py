"""冷启动链路测试：简历文件/文本 → 草稿 → 核对落库，示例数据，画像纠错。

产品主张：新用户从克隆到第一次看到匹配/ATS，不再需要逐字段手抄。
"""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient

from jobhater import config
from jobhater.api import create_app
from jobhater.services.resume_import import (
    ResumeImportError,
    extract_text,
    parse_resume_text,
)

SAMPLE_RESUME = """李明阳
电话：13812345678 | 邮箱：limingyang@example.com
求职意向：后端开发工程师

教育背景
2023.09 - 2027.06　星辰大学　计算机科学与技术　本科

实习经历
2025.06 - 2025.09　云图智能科技　后端开发实习生
- 负责推荐服务日志管线的重构，日志吞吐从 2万/s 提升到 8万/s
- 使用 Go 编写数据同步组件，覆盖 3 条核心链路

项目经历
校园活动平台「聚场」（核心开发）　2024.03-2024.12
- 技术栈：Spring Boot、MySQL、Redis、Docker

专业技能
- 熟悉 Java/Golang，了解 Python
- 熟悉 MySQL、Redis 的使用与调优

荣誉奖项
- 国家励志奖学金（2024）
"""


@pytest.fixture()
def client(tmp_path, monkeypatch):
    config.set_data_dir(tmp_path)
    app = create_app()
    with TestClient(app) as c:
        yield c
    config.set_data_dir(None)


# ---------- 解析器单元 ----------


def test_parse_chinese_resume_full_chain():
    r = parse_resume_text(SAMPLE_RESUME)
    d, w = r["draft"], r["warnings"]
    assert d["basics"]["name"] == "李明阳"
    assert d["basics"]["email"] == "limingyang@example.com"
    assert d["basics"]["phone"] == "13812345678"
    assert d["basics"]["label"] == "后端开发工程师"
    assert d["education"] == [{
        "institution": "星辰大学", "studyType": "本科", "area": "计算机科学与技术",
        "startDate": "2023-09", "endDate": "2027-06",
    }]
    assert len(d["work"]) == 1
    assert d["work"][0]["name"] == "云图智能科技"
    assert d["work"][0]["position"] == "后端开发实习生"
    assert d["work"][0]["startDate"] == "2025-06" and d["work"][0]["endDate"] == "2025-09"
    assert "8万/s" in d["work"][0]["summary"]
    assert d["projects"][0]["name"] == "校园活动平台「聚场」"
    assert d["projects"][0]["role"] == "核心开发"
    names = [s["name"] for s in d["skills"]]
    assert "Java" in names and "Golang" in names and "Python" in names and "Redis" in names
    assert "Redis 的使用与调优" not in names
    assert d["awards"][0]["title"].startswith("国家励志奖学金")
    assert not any("没识别到" in x or "没能识别" in x for x in w)


def test_parse_missing_sections_produce_honest_warnings():
    r = parse_resume_text("王小明\n13800000000\n只是一段没有结构的介绍文本。")
    d = r["draft"]
    assert d["basics"]["name"] == "王小明"
    assert d["basics"]["phone"] == "13800000000"
    warns = " ".join(r["warnings"])
    assert "教育背景" in warns and "技能" in warns


def test_parse_current_job_marker():
    text = "实习经历\n2025.06 - 至今　某公司　开发实习生\n- 做事"
    d = parse_resume_text(text)["draft"]
    assert d["work"][0]["endDate"] == ""


# ---------- 文本抽取 ----------


def test_extract_txt_utf8_and_gb18030(tmp_path):
    assert extract_text("a.txt", SAMPLE_RESUME.encode("utf-8"))[0][:3] == "李明阳"
    gbk_text = "姓名：张三\n教育背景\n2020-2024　南方大学　本科".encode("gb18030")
    assert "南方大学" in extract_text("b.txt", gbk_text)[0]


def test_extract_docx(tmp_path):
    import docx

    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("李明阳")
    doc.add_paragraph("教育背景")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "2021-2025"
    table.rows[0].cells[1].text = "北方大学 本科 计算机"
    doc.save(buf)
    text, kind = extract_text("resume.docx", buf.getvalue())
    assert kind == "docx"
    assert "李明阳" in text and "北方大学" in text


def _make_pdf(content_text: str) -> bytes:
    """程序化构造最小单页 PDF（xref 偏移精确计算），pypdf 可抽取其中的文本。"""
    y = 700
    ops = []
    for ln in content_text.splitlines():
        escaped = ln.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        ops.append(f"BT /F1 12 Tf 72 {y} Td ({escaped}) Tj ET")
        y -= 18
    content = "\n".join(ops).encode("latin-1", errors="replace")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode()
    return bytes(out)


def test_extract_pdf_with_text_layer():
    pdf_text = (
        "Ming Yang\n"
        "Email: ming@example.com | Phone: 13812345678\n"
        "Education\n"
        "2023.09 - 2027.06 Northstar University, Bachelor of Computer Science\n"
        "Skills\n"
        "Java, Golang, Python, MySQL, Redis, Docker"
    )
    text, kind = extract_text("resume.pdf", _make_pdf(pdf_text))
    assert kind == "pdf"
    assert "ming@example.com" in text and "Northstar University" in text
    # PDF 文本继续走解析器：整链可用
    draft = parse_resume_text(text)["draft"]
    assert draft["basics"]["email"] == "ming@example.com"
    assert draft["education"][0]["institution"] == "Northstar University"


def test_extract_pdf_scanned_fails_closed():
    with pytest.raises(ResumeImportError, match="扫描件|没有可提取"):
        extract_text("scan.pdf", _make_pdf(""))


def test_extract_unsupported_and_oversize():
    with pytest.raises(ResumeImportError, match="不支持的文件类型"):
        extract_text("a.exe", b"xx")
    with pytest.raises(ResumeImportError, match="上限"):
        extract_text("a.txt", b"x" * (10 * 1024 * 1024 + 1))


def test_extract_json_passthrough():
    payload = json.dumps({"basics": {"name": "李明阳"}, "skills": [{"name": "Go"}]})
    text, kind = extract_text("r.json", payload.encode("utf-8"))
    assert kind == "json" and json.loads(text)["basics"]["name"] == "李明阳"


# ---------- API：上传/粘贴/AI/落库 ----------


def test_api_parse_resume_endpoints(client):
    # 文件上传（multipart）
    r = client.post(
        "/api/onboarding/parse-resume",
        files={"file": ("resume.md", SAMPLE_RESUME.encode("utf-8"), "text/markdown")},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["source_kind"] == "md"
    assert body["draft"]["basics"]["name"] == "李明阳"
    assert "text_preview" in body and "ai_available" in body
    # 粘贴
    r2 = client.post("/api/onboarding/parse-resume-text", json={"text": SAMPLE_RESUME})
    assert r2.status_code == 200 and r2.json()["draft"]["work"][0]["name"] == "云图智能科技"
    # 上传坏类型 → 422 人话报错
    r3 = client.post(
        "/api/onboarding/parse-resume",
        files={"file": ("x.exe", b"zz", "application/octet-stream")},
    )
    assert r3.status_code == 422 and "不支持的文件类型" in r3.text
    # 空 → 422
    assert client.post("/api/onboarding/parse-resume-text", json={"text": " "}).status_code == 422


def test_api_ai_parse_local_mode(client):
    """本地模式：executed=False 是正常语义（不是错误），前端走启发式结果。"""
    r = client.post("/api/onboarding/ai-parse", json={"text": SAMPLE_RESUME, "ack_egress": False})
    assert r.status_code == 200
    assert r.json()["executed"] is False and r.json()["reason"] == "local_mode"


def test_cold_start_journey_parse_review_import(client):
    """新用户主链：上传→核对（此处直接信任草稿）→建画像→导入→幂等。"""
    draft = client.post(
        "/api/onboarding/parse-resume",
        files={"file": ("resume.txt", SAMPLE_RESUME.encode("utf-8"), "text/plain")},
    ).json()["draft"]

    basics = draft["basics"]
    profile = client.post("/api/profiles", json={
        "display_name": basics["name"] or "未命名",
        "headline": basics["label"] or None,
        "phone": basics["phone"] or None,
        "email": basics["email"] or None,
    }).json()
    client.patch(f"/api/profiles/{profile['id']}", json={"summary": basics["summary"]})
    counts = client.post(f"/api/profiles/{profile['id']}/import/json-resume", json=draft).json()
    assert counts["educations"] == 1 and counts["experiences"] == 1
    assert counts["skills"] >= 4 and counts["projects"] == 1
    # 幂等：重复导入同名实体全部跳过
    counts2 = client.post(f"/api/profiles/{profile['id']}/import/json-resume", json=draft).json()
    assert counts2 == {"skills": 0, "experiences": 0, "educations": 0, "projects": 0,
                       "headline_filled": 0}
    # 画像核对视图能拿到全部导入物
    view = client.get(f"/api/profiles/{profile['id']}").json()
    assert view["profile"]["phone"] == "13812345678"
    assert len(view["skills"]) >= 4 and len(view["educations"]) == 1


def test_api_profile_patch_and_delete_children(client):
    pid = client.post("/api/profiles", json={"display_name": "测试"}).json()["id"]
    # patch：display_name/summary 可改；display_name 拒绝置空
    assert client.patch(f"/api/profiles/{pid}", json={"summary": "一句话"}).json()["summary"] == "一句话"
    assert client.patch(f"/api/profiles/{pid}", json={"display_name": "新名字"}).json()["display_name"] == "新名字"
    assert client.patch(f"/api/profiles/{pid}", json={"display_name": "  "}).status_code == 422
    # 子实体删除：先加再删，删除不存在的 → 404
    edu = client.post(f"/api/profiles/{pid}/educations",
                      json={"school": "某大学", "degree": "本科"}).json()
    assert client.delete(f"/api/profiles/{pid}/educations/{edu['id']}").status_code == 200
    assert client.get(f"/api/profiles/{pid}").json()["educations"] == []
    assert client.delete(f"/api/profiles/{pid}/educations/{edu['id']}").status_code == 404
    exp = client.post(f"/api/profiles/{pid}/experiences",
                      json={"employer": "某公司", "title": "实习生"}).json()
    assert client.delete(f"/api/profiles/{pid}/experiences/{exp['id']}").status_code == 200
    prj = client.post(f"/api/profiles/{pid}/projects", json={"name": "某项目"}).json()
    assert client.delete(f"/api/profiles/{pid}/projects/{prj['id']}").status_code == 200


# ---------- 示例数据 ----------


def test_demo_seed_clear_roundtrip(client):
    s1 = client.post("/api/demo/seed").json()
    assert s1["added"] == 8, s1
    assert client.get("/api/demo/status").json()["demo_jobs"] == 8
    # 幂等：再种一遍全部去重
    s2 = client.post("/api/demo/seed").json()
    assert s2["added"] == 0 and s2["deduped"] == 8
    # 清除 → 干净
    c1 = client.post("/api/demo/clear").json()
    assert c1 == {"deleted": 8, "kept": 0}
    assert client.get("/api/demo/status").json()["demo_jobs"] == 0


def test_demo_clear_keeps_applied_jobs(client):
    client.post("/api/demo/seed")
    jobs = client.get("/api/jobs", params={"limit": 500}).json()["items"]
    demo_job = next(j for j in jobs if j["source_id"] == "demo")
    pid = client.post("/api/profiles", json={"display_name": "测试"}).json()["id"]
    app_id = client.post("/api/applications",
                         json={"job_id": demo_job["id"], "profile_id": pid}).json()["id"]
    r = client.post("/api/demo/clear").json()
    assert r["deleted"] == 7 and r["kept"] == 1
    # RESTRICT 边界：被投递引用的岗位仍在
    assert client.get(f"/api/jobs/{demo_job['id']}").status_code == 200
    assert client.get(f"/api/applications/{app_id}").json()["job_id"] == demo_job["id"]


def test_ai_disclosure_registry_covers_resume_parse(client):
    from jobhater.db import connect
    from jobhater.services.ai import EGRESS_DISCLOSURES, AIService

    assert "resume_parse" in EGRESS_DISCLOSURES
    # 经服务层验证任务目录完整（本地模式 available=False 如实呈现）
    real_con = connect()
    try:
        tasks = {t["task"]: t for t in AIService(real_con).list_tasks()}
    finally:
        real_con.close()
    assert tasks["resume_parse"]["label"] == "简历 AI 精解析"
    assert tasks["resume_parse"]["available"] is False  # 测试环境未配置 provider
